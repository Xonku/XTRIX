"""Report exports: CSV, Excel, standalone HTML (TZ 2.8 item 8 / 3.11 item 9)."""
from __future__ import annotations

import csv
import io
import html
from datetime import datetime, timezone

from .models import ScanResult

COLUMNS = [
    ("endpoint", "Service"), ("days_left", "Days Left"), ("status", "Status"),
    ("risk_score", "Risk Score"), ("risk_level", "Risk Level"),
    ("issuer_cn", "Issuer"), ("subject_cn", "Certificate"),
    ("not_after", "Expiration"), ("chain_valid", "Chain OK"),
    ("hostname_matches", "DNS Match"), ("self_signed", "Self-Signed"),
    ("criticality", "Criticality"), ("owner", "Owner"), ("error", "Error"),
]


def _row(res: ScanResult) -> dict:
    cert = res.cert
    return {
        "endpoint": res.endpoint,
        "days_left": res.days_left,
        "status": res.status,
        "risk_score": res.risk_score,
        "risk_level": res.risk_level,
        "issuer_cn": cert.issuer_cn if cert else None,
        "subject_cn": cert.subject_cn if cert else None,
        "not_after": cert.not_after.strftime("%Y-%m-%d") if cert and cert.not_after else None,
        "chain_valid": res.chain_valid,
        "hostname_matches": res.hostname_matches,
        "self_signed": bool(cert.self_signed) if cert else None,
        "criticality": res.criticality,
        "owner": res.owner,
        "error": res.error,
    }


def to_csv(results: list[ScanResult]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow([title for _, title in COLUMNS])
    for res in results:
        d = _row(res)
        writer.writerow([d[key] for key, _ in COLUMNS])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(results: list[ScanResult]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Certificate Radar"

    header_fill = PatternFill("solid", fgColor="1F3A5F")
    status_fills = {
        "Expired": "F8D7DA", "Critical": "F8D7DA", "Warning": "FFF3CD",
        "Information": "CCE5FF", "OK": "D4EDDA",
    }
    for col, (_, title) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=col, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        ws.column_dimensions[get_column_letter(col)].width = 16

    for i, res in enumerate(results, 2):
        d = _row(res)
        for col, (key, _) in enumerate(COLUMNS, 1):
            c = ws.cell(row=i, column=col, value=d[key])
            if key == "status" and d[key] in status_fills:
                c.fill = PatternFill("solid", fgColor=status_fills[d[key]])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_html(results: list[ScanResult]) -> bytes:
    """Self-contained HTML report (open without any server)."""
    esc = html.escape
    rows = []
    for res in results:
        d = _row(res)
        badge = {"Expired": "crit", "Critical": "crit", "Warning": "warn",
                 "Information": "info", "OK": "ok"}.get(d["status"], "muted")
        chain = {True: "OK", False: "BROKEN", None: "n/a"}.get(res.chain_valid, "n/a")
        dns = {True: "OK", False: "MISMATCH", None: "n/a"}.get(res.hostname_matches, "n/a")
        rows.append(
            f"<tr><td>{esc(d['endpoint'])}</td>"
            f"<td><span class='b {badge}'>{esc(str(d['status']))}</span></td>"
            f"<td>{d['risk_score'] if d['risk_score'] else '-'}</td>"
            f"<td>{d['days_left'] if d['days_left'] is not None else '-'}</td>"
            f"<td>{esc(str(d['subject_cn'] or '-'))}</td>"
            f"<td>{esc(str(d['issuer_cn'] or '-'))}</td>"
            f"<td>{esc(str(d['not_after'] or '-'))}</td>"
            f"<td>{chain}</td><td>{dns}</td>"
            f"<td>{esc(str(d['owner'] or '-'))}</td>"
            f"<td class='err'>{esc(str(d['error'] or ''))}</td></tr>")

    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    return f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<title>Certificate Radar Report</title>
<style>
 body{{font-family:'Segoe UI',Arial,sans-serif;margin:2rem;color:#1e293b}}
 h1{{color:#1f3a5f}} .ts{{color:#64748b;margin-bottom:1rem}}
 table{{border-collapse:collapse;width:100%;font-size:.85rem}}
 th,td{{border:1px solid #cbd5e1;padding:.4rem .5rem;text-align:left}}
 th{{background:#1f3a5f;color:#fff}}
 tr:nth-child(even){{background:#f8fafc}}
 .b{{padding:.15rem .5rem;border-radius:1rem;font-size:.75rem;font-weight:600}}
 .crit{{background:#f8d7da;color:#842029}} .warn{{background:#fff3cd;color:#664d03}}
 .info{{background:#cce5ff;color:#084298}} .ok{{background:#d4edda;color:#0f5132}}
 .muted{{background:#e2e8f0;color:#475569}} .err{{color:#b91c1c;font-size:.75rem}}
</style></head><body>
<h1>Certificate Radar — Report</h1>
<div class="ts">Generated: {ts} · {len(results)} service(s)</div>
<table><thead><tr>
<th>Service</th><th>Status</th><th>Risk</th><th>Days</th><th>Certificate</th>
<th>Issuer</th><th>Expiration</th><th>Chain</th><th>DNS</th><th>Owner</th><th>Error</th>
</tr></thead><tbody>
{''.join(rows)}
</tbody></table></body></html>""".encode("utf-8")
